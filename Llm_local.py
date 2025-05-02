import ollama
import time
import  streamlit as st 
import fitz  
from docx import Document
from io import BytesIO
from fpdf import FPDF

def get_response_from_mistral(query, context):
    prompt_text = f"""
    "Tú eres un asistente para tareas de respuesta a preguntas. "
    "Usa los siguientes fragmentos de contexto recuperado para responder la pregunta. "
    "Si el contexto está vacío o no contiene información relevante, responde: 'Disculpa, no tengo información para responder esa pregunta'. "
    "Si el contexto es válido, responde la pregunta usando un mínimo de 2 oraciones y un máximo de 4, manteniendo la respuesta clara y concisa. "
    "No inventes ni asumas nada que no esté explícitamente en el contexto."
    "\n\n"

    Usa solo este contexto:
    {context}
    **IMPORTANTE***
    "Ojo siempre que tu contexto es vacio, tu respuesta debe ser : Disculpa, no tengo información para responder esa pregunta"
    **
    Y Responde esta pregunta:
    {query}
    """

    respuesta = ollama.chat(
        model="mistral",
        messages=[
            {"role": "system", "content": "Eres un asistente especializado en análisis de datos."},
            {"role": "user", "content": prompt_text}
        ]
    )

    respuesta_texto = respuesta["message"]["content"] 
    
    for word in respuesta_texto.split():  
        yield word + " "
        time.sleep(0.05)

def generarPages():
    with st.sidebar:
        st.page_link("chatbox_v1.py", label="Inicio", icon="🏠")
        st.page_link("pages/resumen_word.py", label="Informe de PDF y Word", icon="📄")
        st.page_link("pages/insertardocumentos.py", label="Documentos a vector", icon="🛢️")

def informes_mistral(context):
    prompt_text = f"""
        **Atención**: No generes una historia o narrativa, tu tarea es realizar un análisis detallado y preciso del documento legal. No se requiere creatividad, solo precisión.
        Eres un asistente experto en procesamiento y análisis de documentos. Tu tarea es leer y comprender el contenido proporcionado y generar un informe extenso, detallado y bien estructurado. 
        El informe debe incluir las siguientes secciones:

        1. **Resumen General**: Proporciona un resumen completo y detallado de todo el contenido del documento. Incluye los aspectos más relevantes, pero sin dejar de lado detalles importantes.
        
        2. **Puntos Clave**: Enumera los puntos más importantes del documento, resaltando las ideas principales y los aspectos críticos que se abordan.

        3. **Análisis Crítico**: Realiza un análisis en profundidad sobre el contenido del documento. Comenta sobre su calidad, lógica, coherencia, posibles fallos, aspectos positivos, y cualquier otro elemento que pueda ser relevante. 

        4. **Recomendaciones**: Proporciona sugerencias o recomendaciones para mejorar el contenido. Si el documento se trata de un informe técnico, análisis de datos, o investigación, incluye sugerencias de cómo se podría mejorar la interpretación de los datos, el análisis o la presentación.

        5. **Conclusiones**: Finaliza con una sección de conclusiones que recapitule los puntos clave del análisis y del documento en general, además de una visión global de las implicaciones del contenido.

        6: **En caso de**: En caso de que el contenido sea acerca de un decreto legislativo o algo acera de una ley incluye un seccion donde hables lo mas importante de todos los articulos y menciones cuales son los mas relevantes.

        7: **Documentos Analizados**: Menciona el nombre de todos los documentos que componen el contenido analizado. Si hay más de uno, asegúrate de listarlos todos y dejar claro que el análisis se basa en todos ellos en conjunto.

        Siempre deberás comenzar el informe con los nombres de los archivos que componen el contenido.
        Recuerda que siempre debes mantener la estructura que te mande

        Contenido del documento:
        {context}

        Utiliza un estilo claro y profesional en todo momento, y asegúrate de que cada sección esté claramente diferenciada. Tu informe debe ser extenso y abarcativo, no debe ser corto ni vago.
        recuerda siempre reponder en español
    """
    respuesta = ollama.chat(
        model="mistral",
        messages=[
            {"role": "system", "content": "Eres un asistente especializado en análisis detallado de documentos."},
            {"role": "user", "content": prompt_text}
        ]
    )

    respuesta_texto = respuesta["message"]["content"] 
    
    for word in respuesta_texto.split():  
        yield word + " "
        time.sleep(0.05)

def extraer_texto(archivo):
    if archivo.name.endswith(".pdf"):
        texto = ""
        with fitz.open(stream=archivo.read(), filetype="pdf") as doc:
            for page in doc:
                texto += page.get_text()
        return texto
    elif archivo.name.endswith(".txt"):
        return archivo.read().decode("utf-8")
    else:
        return ""

def extraer_texto_word(file):
    texto = ""
    doc = Document(file)
    for para in doc.paragraphs:
        texto += para.text + "\n"
    return texto

def generar_docx(texto):
    doc = Document()
    doc.add_heading("Resumen generado por IA", 0)
    for parrafo in texto.split("\n"):
        doc.add_paragraph(parrafo)
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

def generar_pdf(texto):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.set_font("Arial", size=12)
    for linea in texto.split("\n"):
        pdf.multi_cell(0, 10, linea)
    return bytes(pdf.output(dest='S').encode('latin-1'))
